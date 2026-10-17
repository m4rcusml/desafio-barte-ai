#!/usr/bin/env python3
"""
Massa de teste para o case Tarvo (Desafio Barte AI).

Gera, no MESMO formato do data pack de setembro/2026, arquivos de entrada com
lojistas inventados e pegadinhas plantadas, e o gabarito esperado no formato
de saida/.

O gabarito NAO e calculado relendo os arquivos gerados: ele sai da "verdade"
que o proprio gerador planta (cada linha sabe de que lojista e, qual taxa foi
contratada e qual pegadinha carrega). Assim a comparacao com a solucao do time
e independente do parser.

Uso:
    python3 gerar_massa.py                 # gera entrada/ + gabarito/ em ./
    python3 gerar_massa.py --saida DIR     # em outra pasta
    python3 gerar_massa.py --so-entradas   # so os arquivos de entrada

Deterministico: mesma SEED -> mesmos arquivos.
"""
import argparse
import base64
import csv
import datetime as dt
import io
import os
import random
import shutil
import unicodedata
from decimal import Decimal as D, ROUND_DOWN, ROUND_HALF_UP

import openpyxl
from openpyxl.comments import Comment

SEED = 20260917
AQUI = os.path.dirname(os.path.abspath(__file__))
C2 = D("0.01")


# ----------------------------------------------------------------- utilidades
def q2(x):
    return x.quantize(C2, ROUND_HALF_UP)


def trunc2(x):
    """Orbe trunca a taxa em centavos (em direcao ao zero), como no arquivo real."""
    return x.quantize(C2, ROUND_DOWN)


def br(x, milhar=False):
    s = f"{x:,.2f}"
    s = s.replace(",", "X").replace(".", ",").replace("X", ".")
    return s if milhar else s.replace(".", "")


def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def cnpj_dv(base12):
    def dv(nums, pesos):
        r = sum(int(n) * p for n, p in zip(nums, pesos)) % 11
        return "0" if r < 2 else str(11 - r)

    d1 = dv(base12, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    d2 = dv(base12 + d1, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    return base12 + d1 + d2


def fmt_cnpj(c):
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}"


def d(dia, mes=9):
    return dt.date(2026, mes, dia)


def dbr(x):
    return x.strftime("%d/%m/%Y")


FERIADOS = {d(7)}
DIAS_UTEIS = [d(i) for i in range(1, 31) if d(i).weekday() < 5 and d(i) not in FERIADOS]


def prox_util(x):
    y = x + dt.timedelta(days=1)
    while y.weekday() >= 5 or y in FERIADOS:
        y += dt.timedelta(days=1)
    return y


# ------------------------------------------------------------------- mundo
STD = dict(deb=D("1.49"), avista=D("2.69"), p2_6=D("3.19"), p7_12=D("3.79"), pix=D("0.89"), antec=D("2.10"))
ANTIGO_2025 = dict(deb=D("1.69"), avista=D("2.99"), parc=D("3.49"), pix=D("0.99"))


def tx(**kw):
    t = dict(STD)
    t.update({k: D(v) for k, v in kw.items()})
    return t


# cnpj_fmt: como o CNPJ aparece na planilha de taxas ('fmt', 'digitos', 'int' = numero sem zero a esquerda)
# ec_tab: se o codigo EC aparece na planilha vigente
# aba: 'vigente', 'antigo' (so na aba ANTIGO 2025) ou 'ambas'
LOJISTAS = [
    dict(k="atlantica", tabela="Atlântica Intercâmbio Ltda", razao="Atlântica Intercâmbio Ltda", cnpj_fmt="fmt",
         ec_tab=True, orbe=["ATLANTICA INTERCAMBIO", "ATLANTICA INTERC LTDA"], tx=tx(avista="2.59"),
         spread=D("0.015"), vig="01/01/2026", obs=None, aba="ambas", checkout=True, remessa=True, ponte_cnpj="fmt"),
    dict(k="pontofinal", tabela="Escola Ponto Final de Idiomas Ltda", razao="Escola Ponto Final de Idiomas Ltda",
         cnpj_fmt="fmt", ec_tab=True, orbe=["PONTO FINAL IDIOMAS", "ESCOLA PONTO FINAL IDIOMAS"],
         tx=tx(avista="2.59"), spread=None, vig="até 14/09/2026", obs="renegociado - ver linha de baixo",
         aba="ambas", checkout=True, remessa=False),
    dict(k="rumonorte", tabela="Rumo Norte Viagens Estudantis Ltda", razao="Rumo Norte Viagens Estudantis Ltda",
         cnpj_fmt="digitos", ec_tab=True, orbe=["RUMO NORTE VIAGENS"],
         tx=tx(avista="2.59", p2_6="3.09"), spread=None, vig="01/01/2026", obs=None, aba="vigente",
         checkout=True, remessa=False, taxa_texto=True),
    dict(k="wombat", tabela="Wombat Experience", razao="Wombat Experience Turismo Ltda", cnpj_fmt="digitos",
         ec_tab=False, orbe=["WOMBAT EXPERIENCE TURISMO"], tx=tx(avista="2.59"), spread=None,
         vig="01/01/2026", obs="EC pendente", aba="ambas", checkout=True, remessa=True, ponte_cnpj="digitos"),
    dict(k="caminhos", tabela="Caminhos Intercâmbio", razao="Caminhos Intercâmbio Ltda", cnpj_fmt="fmt",
         ec_tab=True, orbe=["CAMINHOS INTERCAMBIO", "Caminhos", "CAMINHOS INTERC LTDA"],
         tx=tx(avista="2.49", p2_6="2.99"), spread=D("0.012"), vig="01/01/2026",
         obs="spread negociado em jul/26", aba="ambas", checkout=True, remessa=True, ponte_cnpj="fmt"),
    dict(k="fiorde", tabela="Fiorde Idiomas Online Ltda", razao="Fiorde Idiomas Online Ltda", cnpj_fmt="int",
         ec_tab=True, orbe=["FIORDE IDIOMAS ONLINE"], tx=tx(pix="0.79"), spread=None, vig="01/01/2026",
         obs=None, aba="vigente", checkout=True, remessa=False),
    dict(k="cordilheira", tabela="Cordilheira Expedições Ltda", razao="Cordilheira Expedições Ltda",
         cnpj_fmt="digitos", ec_tab=True, orbe=["CORDILHEIRA EXPEDICOES"], tx=tx(deb="1.39"), spread=None,
         vig="01/01/2026", obs=None, aba="vigente", checkout=True, remessa=False),
    dict(k="casabrasil", tabela="Casa Brasil Estudantes Ltda", razao="Casa Brasil Estudantes Ltda",
         cnpj_fmt="digitos", ec_tab=True, orbe=["CASA BRASIL ESTUDANTES", "CASA BRASIL ESTUD"], tx=tx(),
         spread=D("0.015"), vig="01/01/2026", obs="cliente estratégico - ver c/ comercial", aba="ambas",
         checkout=True, remessa=True, ponte_cnpj="digitos",
         comentario="conversar com Marcos (comercial) sobre condição especial"),
    dict(k="campusmundi", tabela="Campus Mundi Ltda", razao="Campus Mundi Ltda", cnpj_fmt="int", ec_tab=True,
         orbe=["CAMPUS MUNDI"], tx=tx(p2_6="2.99"), spread=None, vig="01/01/2026", obs=None, aba="vigente",
         checkout=True, remessa=False),
    dict(k="trilhasul", tabela="Trilha Sul Turismo Ltda", razao="Trilha Sul Turismo Ltda", cnpj_fmt="digitos",
         ec_tab=True, orbe=["TRILHA SUL TURISMO"], tx=tx(), spread=None, vig="01/01/2026", obs=None,
         aba="ambas", checkout=True, remessa=False),
    dict(k="boasvindas", tabela="Boas-Vindas Exchange Ltda", razao="Boas-Vindas Exchange Ltda",
         cnpj_fmt="digitos", ec_tab=True, orbe=["BOAS-VINDAS EXCHANGE", "BOAS VINDAS EXCHANGE"], tx=tx(),
         spread="padrão", vig="01/01/2026", obs=None, aba="vigente", checkout=True, remessa=True,
         ponte_cnpj="digitos"),
    dict(k="lisboeta", tabela="Lisboeta Cursos e Viagens Ltda", razao="Lisboeta Cursos e Viagens Ltda",
         cnpj_fmt="fmt", ec_tab=True, orbe=["LISBOETA CURSOS E VIAGENS", "LISBOETA CURSOS VIAGENS"], tx=tx(),
         spread=D("0.016"), vig="01/06/2026", obs=None, aba="ambas", checkout=True, remessa=True,
         ponte_cnpj="fmt"),
    dict(k="estrelaviagens", tabela="Estrela Viagens", razao="Estrela Viagens e Turismo Ltda",
         cnpj_fmt="digitos", ec_tab=False, orbe=["ESTRELA VIAGENS TURISMO"], tx=tx(), spread=None,
         vig="01/01/2026", obs=None, aba="ambas", checkout=True, remessa=False),
    dict(k="novaestrela", tabela="Nova Estrela Intercâmbio", razao="Nova Estrela Intercâmbio Ltda",
         cnpj_fmt="digitos", ec_tab=False, orbe=["ESTRELA INTERCAMBIO"], tx=tx(avista="2.39", p2_6="2.89"),
         spread=D("0.014"), vig="01/06/2026", obs=None, aba="ambas", checkout=True, remessa=True,
         ponte_cnpj="fmt"),
    dict(k="lumen", tabela=None, razao="Lumen Idiomas Ltda", cnpj_fmt="digitos", ec_tab=False,
         orbe=["LUMEN IDIOMAS"], tx=tx(), spread=None, vig=None, obs=None, aba="antigo", checkout=True,
         remessa=False),
    dict(k="globoandante", tabela="Globo Andante Viagens Ltda", razao="Globo Andante Viagens Ltda",
         cnpj_fmt="digitos", ec_tab=True, orbe=[], tx=tx(), spread=None, vig=None, obs="INATIVO desde jul/26",
         aba="vigente", checkout=False, remessa=False),
    dict(k="polar", tabela="Polar Educação Internacional Ltda", razao="Polar Educação Internacional Ltda",
         cnpj_fmt="digitos", ec_tab=False, orbe=[], tx=None, spread=D("0.0125"), vig="01/01/2026",
         obs="só remessa", aba="vigente", checkout=False, remessa=True, ponte_cnpj="fmt"),
    dict(k="rotalivre", tabela="Rota Livre Estudantil Ltda", razao="Rota Livre Estudantil Ltda",
         cnpj_fmt="digitos", ec_tab=False, orbe=[], tx=None, spread=None, vig="01/01/2026", obs="só remessa",
         aba="vigente", checkout=False, remessa=True, ponte_cnpj="digitos"),
]
L = {x["k"]: x for x in LOJISTAS}

TARVO_CNPJ = "84626315000139"
PONTE_CNPJ = "75374179000146"
PONTE_RAZAO = "Ponte Câmbio S.A."
SPREAD_PADRAO = D("0.015")


def montar_cadastro(rng):
    for x in LOJISTAS:
        zero = x["cnpj_fmt"] == "int"  # forca zero a esquerda para ele "sumir" no Excel
        x["cnpj"] = cnpj_dv(("0" if zero else str(rng.randint(1, 9)))
                            + "".join(str(rng.randint(0, 9)) for _ in range(7)) + "0001")
        x["ec"] = str(rng.randint(20, 39)) + "".join(str(rng.randint(0, 9)) for _ in range(8))
    L["casabrasil"]["cnpj"] = cnpj_dv("00" + L["casabrasil"]["cnpj"][2:12])


def spread_contratado(x):
    s = x["spread"]
    if s is None or s == "padrão":
        return SPREAD_PADRAO
    return s


def taxa_contrato(x, tipo, mod, parc, data_venda):
    t = x["tx"]
    if tipo in ("ANTECIPACAO", "REVERSAO ANTECIPACAO"):
        return t["antec"]
    if mod == "DEBITO":
        return t["deb"]
    if mod == "PIX":
        return t["pix"]
    if mod == "CREDITO A VISTA":
        if x["k"] == "pontofinal" and data_venda >= d(15):
            return D("2.29")
        return t["avista"]
    return t["p2_6"] if parc <= 6 else t["p7_12"]


# -------------------------------------------------------------- Orbe (01/06)
BANDEIRAS = ["VISA", "MASTERCARD", "ELO", "AMEX", "HIPERCARD"]


class Orbe:
    def __init__(self, rng):
        self.rng = rng
        self.rows = []

    def _nova(self, x, data, tipo, mod, parc, band, bruto, taxa, ref="", data_orig=None, arq="01",
              data_venda=None, tags=()):
        vt = trunc2(bruto * taxa / 100)
        r = dict(id=None, data=data, ec=x["ec"], nome=self.rng.choice(x["orbe"]), tipo=tipo, mod=mod, parc=parc,
                 band=band, bruto=bruto, taxa=taxa, vt=vt, liq=bruto - vt, ref=ref, data_orig=data_orig, arq=arq,
                 lk=x["k"], data_venda=data_venda or data, tags=list(tags))
        r["ct"] = taxa_contrato(x, tipo, mod, parc, r["data_venda"])
        self.rows.append(r)
        return r

    def venda(self, x, data, mod=None, parc=None, bruto=None, taxa=None, arq="01", tags=()):
        rng = self.rng
        if mod is None:
            mod = rng.choices(["CREDITO PARCELADO", "CREDITO A VISTA", "PIX", "DEBITO"], [50, 30, 10, 10])[0]
        if parc is None:
            parc = rng.randint(2, 12) if mod == "CREDITO PARCELADO" else 1
        if bruto is None:
            faixa = {"CREDITO PARCELADO": (900, 11000), "CREDITO A VISTA": (250, 3200), "PIX": (180, 4200),
                     "DEBITO": (90, 1600)}[mod]
            bruto = D(str(round(rng.uniform(*faixa), 2)))
        band = "" if mod == "PIX" else rng.choice(BANDEIRAS)
        if taxa is None:
            taxa = taxa_contrato(x, "VENDA", mod, parc, data)
        return self._nova(x, data, "VENDA", mod, parc, band, bruto, D(taxa), arq=arq, tags=tags)

    def estorno(self, venda, data, arq="01", tags=()):
        x = L[venda["lk"]]
        band = "" if venda["mod"] == "PIX" else self.rng.choice(BANDEIRAS)  # bandeira nao bate: ruido do export real
        r = self._nova(x, data, "ESTORNO", venda["mod"], venda["parc"], band, -venda["bruto"], venda["taxa"],
                       ref=venda["id_tmp"], data_orig=venda["data"], arq=arq, data_venda=venda["data"], tags=tags)
        venda["estornada"] = True
        return r

    def antecipacao(self, x, data, bruto):
        return self._nova(x, data, "ANTECIPACAO", "", None, "", bruto, x["tx"]["antec"])

    def reversao(self, ant, data, arq="01", tags=()):
        x = L[ant["lk"]]
        r = self._nova(x, data, "REVERSAO ANTECIPACAO", "", None, "", -ant["bruto"], ant["taxa"],
                       ref=ant["id_tmp"], data_orig=ant["data"], arq=arq, data_venda=ant["data"], tags=tags)
        ant["revertida"] = True
        return r


def gerar_orbe(rng):
    o = Orbe(rng)
    # id provisorio para ligar estornos/reversoes; id definitivo sai em ordem cronologica no fim
    def marca(r):
        r["id_tmp"] = id(r)
        return r

    loj_checkout = [x for x in LOJISTAS if x["checkout"]]
    for x in loj_checkout:
        for _ in range(rng.randint(14, 18)):
            marca(o.venda(x, d(rng.randint(1, 29))))

    # ---- pegadinhas plantadas (venda) ----
    pf = L["pontofinal"]
    marca(o.venda(pf, d(14), "CREDITO A VISTA", 1, D("1480.00"), tags=["VIGENCIA_FRONTEIRA"]))
    marca(o.venda(pf, d(15), "CREDITO A VISTA", 1, D("1325.40"), tags=["VIGENCIA_FRONTEIRA"]))
    for dia, valor in [(16, "2210.35"), (19, "985.10"), (23, "1760.00"), (27, "2433.90")]:
        marca(o.venda(pf, d(dia), "CREDITO A VISTA", 1, D(valor), taxa="2.59", tags=["DIVERG_TAXA_ANTIGA"]))
    venda_vig = marca(o.venda(pf, d(10), "CREDITO A VISTA", 1, D("1890.00"), tags=["VIGENCIA_ESTORNO"]))

    rn = L["rumonorte"]
    for dia, p, valor in [(5, 3, "4120.00"), (12, 5, "2870.55"), (21, 4, "6310.80")]:
        marca(o.venda(rn, d(dia), "CREDITO PARCELADO", p, D(valor), taxa="3.19", tags=["DIVERG_TAXA_TEXTO"]))

    fi = L["fiorde"]
    for dia, valor in [(8, "3120.00"), (24, "1876.45")]:
        marca(o.venda(fi, d(dia), "PIX", 1, D(valor), taxa="0.89", tags=["DIVERG_PIX"]))
    for dia, valor in [(11, "1.00"), (26, "3.50")]:
        marca(o.venda(fi, d(dia), "PIX", 1, D(valor), taxa="0.89", tags=["R3_CENTAVO"]))

    co = L["cordilheira"]
    div_cord = [marca(o.venda(co, d(dia), "DEBITO", 1, D(valor), taxa="1.49", tags=["DIVERG_DEBITO"]))
                for dia, valor in [(3, "845.20"), (17, "1210.00"), (22, "640.75")]]

    ts = L["trilhasul"]
    div_trilha = [marca(o.venda(ts, d(dia), "CREDITO A VISTA", 1, D(valor), taxa="2.99", tags=["DIVERG_AVISTA"]))
                  for dia, valor in [(4, "2650.00"), (9, "1150.30"), (25, "3010.00")]]

    cm = L["campusmundi"]
    for dia, p, valor in [(6, 2, "5400.00"), (18, 6, "7725.15")]:
        marca(o.venda(cm, d(dia), "CREDITO PARCELADO", p, D(valor), taxa="3.19", tags=["DIVERG_PARCELADO"]))

    wb = L["wombat"]
    marca(o.venda(wb, d(13), "CREDITO A VISTA", 1, D("2299.00"), taxa="2.69", tags=["DIVERG_CASADO_POR_NOME"]))

    # ---- antecipacoes e reversoes ----
    antecs = []
    for k in ["atlantica", "caminhos", "trilhasul", "lisboeta", "casabrasil", "rumonorte", "wombat",
              "boasvindas", "campusmundi"]:
        antecs.append(marca(o.antecipacao(L[k], d(rng.randint(2, 24)), D(str(round(rng.uniform(15000, 55000), 2))))))
    rng.shuffle(antecs)
    for a in antecs[:2]:
        o.reversao(a, a["data"] + dt.timedelta(days=rng.randint(2, 4)), tags=["REVERSAO_ANTECIPACAO"])
    for a in antecs[2:4]:
        o.reversao(a, d(30), arq="06", tags=["AJUSTE_TARDIO", "REVERSAO_ANTECIPACAO"])

    # ---- estornos ----
    vendas_ok = [r for r in o.rows if r["tipo"] == "VENDA" and not r["tags"] and r["data"] <= d(26)]
    rng.shuffle(vendas_ok)
    o.estorno(venda_vig, d(22), tags=["VIGENCIA_ESTORNO"])
    o.estorno(div_trilha[1], d(16), tags=["R4_ESTORNO_DIVERGENTE"])
    usados = 0
    for v in vendas_ok:
        if usados >= 8:
            break
        if v.get("estornada"):
            continue
        o.estorno(v, v["data"] + dt.timedelta(days=rng.randint(2, 4)), tags=["ESTORNO"])
        usados += 1
    # garante estorno de lojista sem contrato e de lojista ambiguo
    for k in ["lumen", "novaestrela"]:
        v = next(r for r in o.rows if r["tipo"] == "VENDA" and r["lk"] == k and not r.get("estornada")
                 and r["data"] <= d(25))
        o.estorno(v, v["data"] + dt.timedelta(days=3), tags=["ESTORNO"])
    # ajustes tardios (arquivo 06): estornos com data 30/09
    o.estorno(div_cord[0], d(30), arq="06", tags=["AJUSTE_TARDIO", "R4_ESTORNO_DIVERGENTE"])
    usados = 0
    for v in vendas_ok:
        if usados >= 6:
            break
        if v.get("estornada"):
            continue
        o.estorno(v, d(30), arq="06", tags=["AJUSTE_TARDIO", "ESTORNO"])
        usados += 1
    o.venda(L["wombat"], d(30), "CREDITO PARCELADO", 3, D("5120.00"), arq="06",
            tags=["AJUSTE_TARDIO", "VENDA_TARDIA"])

    # ---- ids definitivos (cronologicos, com saltos) ----
    o.rows.sort(key=lambda r: (r["data"], r["arq"], rng.random()))
    n_tx, n_an = 77310412, 77349870
    mapa = {}
    for r in o.rows:
        if r["tipo"] in ("ANTECIPACAO", "REVERSAO ANTECIPACAO"):
            n_an += rng.randint(9, 40)
            r["id"] = f"AN{n_an}"
        else:
            n_tx += rng.randint(9, 60)
            r["id"] = f"TX{n_tx}"
        if "id_tmp" in r:
            mapa[r["id_tmp"]] = r["id"]
    for r in o.rows:
        if r["ref"]:
            r["ref"] = mapa[r["ref"]]
    return o.rows


ORBE_HEADER = ("id_transacao;data_movimento;cod_estabelecimento;nome_estabelecimento;tipo_lancamento;modalidade;"
               "qtd_parcelas;bandeira;valor_bruto;taxa_aplicada_pct;valor_taxa;valor_liquido;id_referencia;"
               "data_venda_original")


def escrever_orbe(rows, caminho, arq, rng):
    sel = [r for r in rows if r["arq"] == arq]
    # dentro do mesmo dia o portal nao ordena por id
    por_dia = {}
    for r in sel:
        por_dia.setdefault(r["data"], []).append(r)
    linhas = [ORBE_HEADER]
    for dia in sorted(por_dia):
        grupo = por_dia[dia]
        rng.shuffle(grupo)
        for r in grupo:
            linhas.append(";".join([
                r["id"], dbr(r["data"]), r["ec"], r["nome"], r["tipo"], r["mod"],
                "" if r["parc"] is None else str(r["parc"]), r["band"], br(r["bruto"]), br(r["taxa"]),
                br(r["vt"]), br(r["liq"]), r["ref"], dbr(r["data_orig"]) if r["data_orig"] else ""]))
    with open(caminho, "w", encoding="ascii", newline="") as f:
        f.write("\r\n".join(linhas) + "\r\n")


# ------------------------------------------------------- tabela de taxas xlsx
def escrever_tabela(caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Taxas vigentes"
    ws["A1"] = "TARVO - TABELA DE TAXAS CONTRATADAS (atualizar sempre que o comercial fechar!!)"
    ws["A2"] = "última atualização: 03/09 - Bia"
    ws["D3"] = "CARTÃO / PIX (MDR)"
    ws["J3"] = "REMESSA"
    cab = ["Lojista", "CNPJ", "EC Orbe", "Débito", "Crédito à vista", "Créd. 2-6x", "Créd. 7-12x", "Pix",
           "Antecipação", "Spread remessa", "Vigência", "Obs"]
    for i, c in enumerate(cab, 1):
        ws.cell(4, i, c)

    def cnpj_cel(x):
        if x["cnpj_fmt"] == "fmt":
            return fmt_cnpj(x["cnpj"])
        if x["cnpj_fmt"] == "int":
            return int(x["cnpj"])
        return x["cnpj"]

    def pct(v, texto):
        if v is None:
            return None
        return (f"{v:.2f}".replace(".", ",") + "%") if texto else float(v / 100)

    ordem = ["pontofinal", "PF2", "estrelaviagens", "wombat", "atlantica", "lisboeta", "cordilheira", "casabrasil",
             "fiorde", "caminhos", "polar", "trilhasul", "rumonorte", "boasvindas", "campusmundi", "rotalivre",
             "novaestrela", "globoandante"]
    linha = 5
    linhas_tab = {}
    for k in ordem:
        if k == "PF2":
            x = dict(L["pontofinal"], tabela="Ponto Final Idiomas (NOVA)", tx=dict(L["pontofinal"]["tx"],
                     avista=D("2.29")), vig="a partir de 15/09/2026", obs="só mudou o crédito à vista")
        else:
            x = L[k]
        texto = x.get("taxa_texto", False)
        t = x["tx"]
        sp = x["spread"]
        vals = [x["tabela"], cnpj_cel(x), x["ec"] if x["ec_tab"] else None]
        if t is None:
            vals += [None] * 6
        else:
            vals += [pct(t[c], texto) for c in ["deb", "avista", "p2_6", "p7_12", "pix", "antec"]]
        vals += [float(sp) if isinstance(sp, D) else sp, x["vig"], x["obs"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(linha, i, v)
            if isinstance(v, float) and 4 <= i <= 10:
                c.number_format = "0.00%" if i < 10 else "0.0%"
        if x.get("comentario"):
            ws.cell(linha, 1).comment = Comment(x["comentario"], "Bia")
        linhas_tab[k] = linha
        linha += 1
    ws.cell(linha + 1, 1, "* spread em branco = padrão")

    ws2 = wb.create_sheet("ANTIGO 2025 - nao usar")
    for i, c in enumerate(["Lojista", "CNPJ", "Débito", "Crédito à vista", "Créd. parcelado", "Pix"], 1):
        ws2.cell(1, i, c)
    r = 2
    for k in ["caminhos", "atlantica", "pontofinal", "casabrasil", "estrelaviagens", "novaestrela", "lumen",
              "wombat", "lisboeta", "trilhasul"]:
        x = L[k]
        nome = x["razao"]
        vals = [nome, x["cnpj"], float(ANTIGO_2025["deb"] / 100), float(ANTIGO_2025["avista"] / 100),
                float(ANTIGO_2025["parc"] / 100), float(ANTIGO_2025["pix"] / 100)]
        for i, v in enumerate(vals, 1):
            ws2.cell(r, i, v)
        r += 1
    wb.save(caminho)
    return linhas_tab


# --------------------------------------------------------------- Ponte (eml)
MOEDAS = {"USD": D("5.4120"), "EUR": D("6.2530"), "GBP": D("7.1980"), "CAD": D("3.9560"), "AUD": D("3.5540"),
          "NZD": D("3.2470")}


def gerar_ponte(rng):
    plano = [("atlantica", 4), ("wombat", 3), ("caminhos", 4), ("casabrasil", 4), ("boasvindas", 4),
             ("lisboeta", 4), ("novaestrela", 4), ("polar", 4), ("rotalivre", 4)]
    ops = []
    for k, n in plano:
        for _ in range(n):
            ops.append(dict(lk=k, dc=rng.choice(DIAS_UTEIS), tags=[]))
    rng.shuffle(ops)
    # spread negativo (taxa travada D-1) em lisboeta e rotalivre
    for k in ["lisboeta", "rotalivre"]:
        next(o for o in ops if o["lk"] == k)["neg"] = True
    ress = [o for o in ops if not o.get("neg")]
    for o in rng.sample(ress, 3):
        o["ressalva"] = True
    n = 2
    for o in sorted(ops, key=lambda o: o["dc"]):
        n += rng.randint(1, 4)
        o["num"] = f"PC-0926-{n:04d}"
    for o in ops:
        x = L[o["lk"]]
        moeda = rng.choice(list(MOEDAS))
        ptax = (MOEDAS[moeda] * D(str(round(rng.uniform(0.985, 1.015), 4)))).quantize(D("0.0001"))
        me = D(str(round(rng.uniform(1500, 12000), 2)))
        if o.get("neg"):
            s = D("-0.0155")
            o["tags"].append("SPREAD_NEGATIVO")
            o["obs"] = "Taxa travada no pedido (D-1)"
        else:
            s = D("0.015") if x["k"] == "caminhos" else spread_contratado(x)  # Ponte aplica 1,5% p/ Caminhos
            o["obs"] = None
            if x["k"] == "caminhos":
                o["tags"].append("DIVERG_SPREAD")
            if x["spread"] is None:
                o["tags"].append("SPREAD_EM_BRANCO")
            if x["spread"] == "padrão":
                o["tags"].append("SPREAD_TEXTO_PADRAO")
        taxa = (ptax / (1 - s)).quantize(D("0.0001"), ROUND_HALF_UP)
        valor = q2(me * taxa)
        spread = q2((taxa - ptax) * me)
        est_base = abs(spread) if o.get("neg") else spread
        o.update(moeda=moeda, ptax=ptax, me=me, taxa=taxa, valor=valor, spread=spread,
                 spread_pct=(spread / valor).quantize(D("0.000001"), ROUND_HALF_UP),
                 est=q2(est_base * D(str(round(rng.uniform(0.88, 1.12), 3)))), iof=q2(valor * D("0.0038")),
                 status="APROVADO C/ RESSALVA" if o.get("ressalva") else "APROVADO", dl=prox_util(o["dc"]),
                 dup=False)
        if o.get("ressalva"):
            o["tags"].append("RESSALVA")
        cn = x["cnpj"]
        o["cnpj_txt"] = fmt_cnpj(cn) if x["ponte_cnpj"] == "fmt" else cn
        o["cliente"] = x["razao"].upper()[:20]
    ops.sort(key=lambda o: o["num"])
    # operacao reprocessada: mesma linha repetida logo abaixo
    alvo = next(o for o in ops if o["lk"] == "boasvindas" and not o.get("ressalva"))
    dup = dict(alvo, obs="reprocessada", dup=True, tags=["OPERACAO_DUPLICADA"])
    alvo["tags"] = alvo["tags"] + ["OPERACAO_DUPLICADA_ORIGINAL"]
    ops.insert(ops.index(alvo) + 1, dup)
    return ops


def escrever_ponte_xlsx(ops, caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Operações"
    ws["A1"] = "PONTE CÂMBIO S.A. - Relatório Mensal de Operações por Parceiro"
    ws["A2"] = "Parceiro: TARVO PAGAMENTOS LTDA  |  CNPJ 84.626.315/0001-39"
    ws["A3"] = "Competência: 09/2026  |  Gerado em 02/10/2026 07:12"
    cab = ["Nº Operação", "Data Contratação", "Data Liquidação", "CNPJ Cliente", "Cliente", "Moeda", "Valor ME",
           "Taxa Cliente", "PTAX Ref.", "Valor R$", "Spread R$", "Spread %", "Spread Estimado R$ (D-1)", "IOF R$",
           "Status Compliance", "Observação"]
    for i, c in enumerate(cab, 1):
        ws.cell(5, i, c)
    r = 6
    for o in ops:
        o["linha_xlsx"] = r
        vals = [o["num"], dt.datetime.combine(o["dc"], dt.time()), dt.datetime.combine(o["dl"], dt.time()),
                o["cnpj_txt"], o["cliente"], o["moeda"], float(o["me"]), float(o["taxa"]), float(o["ptax"]),
                float(o["valor"]), float(o["spread"]), float(o["spread_pct"]), float(o["est"]), float(o["iof"]),
                o["status"], o["obs"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            if i in (2, 3):
                c.number_format = "DD/MM/YYYY"
            elif i in (8, 9):
                c.number_format = "0.0000"
            elif i == 12:
                c.number_format = "0.000%"
        r += 1
    ws.cell(r + 1, 1, "Total de operações:")
    ws.cell(r + 1, 2, len(ops))
    ws.cell(r + 2, 1, "Spread total R$:")
    ws.cell(r + 2, 2, float(sum(o["spread"] for o in ops)))
    ws2 = wb.create_sheet("Resumo")
    for i, c in enumerate(["Moeda", "Qtd", "Volume ME", "Spread R$"], 1):
        ws2.cell(1, i, c)
    for j, m in enumerate(["USD", "EUR", "GBP", "CAD", "AUD", "NZD"], 2):
        sel = [o for o in ops if o["moeda"] == m]
        ws2.cell(j, 1, m)
        ws2.cell(j, 2, len(sel))
        ws2.cell(j, 3, float(sum(o["me"] for o in sel)))
        ws2.cell(j, 4, float(sum(o["spread"] for o in sel)))
    bio = io.BytesIO()
    wb.save(bio)
    if caminho:
        with open(caminho, "wb") as f:
            f.write(bio.getvalue())
    return bio.getvalue()


EML_CABECALHO = """From:
 =?utf-8?q?Relat=C3=B3rios_Ponte_C=C3=A2mbio?= <relatorios@pontecambio.com.br>
To: financeiro@tarvo.com.br
Cc: comercial.parceiros@pontecambio.com.br
Date: Fri, 02 Oct 2026 07:15:03 -0300
Subject: =?utf-8?q?Relat=C3=B3rio_mensal_de_opera=C3=A7=C3=B5es?= - TARVO -
 09/2026
MIME-Version: 1.0
Content-Type: multipart/mixed; boundary="===============7107453185626437558=="

--===============7107453185626437558==
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: quoted-printable

Prezados,

Segue em anexo o relat=C3=B3rio mensal de opera=C3=A7=C3=B5es de c=C3=A2mbio =
do parceiro TARVO PAGAMENTOS referente =C3=A0 compet=C3=AAncia 09/2026.

Conforme contrato de parceria, o repasse da participa=C3=A7=C3=A3o de 50% (ci=
nquenta por cento) do spread realizado ser=C3=A1 creditado em 06/10/2026, med=
iante recebimento da nota fiscal de intermedia=C3=A7=C3=A3o.

Lembramos que opera=C3=A7=C3=B5es com status "APROVADO C/ RESSALVA" seguem no=
rmalmente para liquida=C3=A7=C3=A3o.

Em caso de diverg=C3=AAncia, favor responder este e-mail em at=C3=A9 5 dias =
=C3=BAteis.

Atenciosamente,
Backoffice Parceiros - Ponte C=C3=A2mbio S.A. - CNPJ 75.374.179/0001-46
Este =C3=A9 um e-mail autom=C3=A1tico.

--===============7107453185626437558==
Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet
Content-Transfer-Encoding: base64
Content-Disposition: attachment;
 filename="PonteCambio_Relatorio_TARVO_092026.xlsx"
MIME-Version: 1.0

"""


def escrever_eml(xlsx_bytes, caminho):
    b64 = base64.encodebytes(xlsx_bytes).decode("ascii")  # linhas de 76 chars, como o original
    texto = EML_CABECALHO + b64 + "\n--===============7107453185626437558==--\n"
    with open(caminho, "w", encoding="ascii", newline="\n") as f:
        f.write(texto)


# ---------------------------------------------------------- extrato Aurora
def nome_aurora(x):
    if x["k"] == "cordilheira":
        return "CORDILHEIRA EXPEDICÕES LTD"  # o banco real mantem o Õ e tira o Ç
    return sem_acento(x["razao"]).upper()[:26]


def gerar_aurora(rng):
    ev = []  # (data, descricao, documento, valor, categoria, lojista)

    def add(dia, desc, valor, cat, doc="", lk=None):
        ev.append(dict(data=d(dia), desc=desc, doc=doc, valor=D(valor), cat=cat, lk=lk))

    def ted_orbe(dia):
        add(dia, "TED RECEBIDA 99.876.669/0001-09 ORBE ADQUIRENCIA SA", str(round(rng.uniform(18000, 62000), 2)),
            "ORBE_TED", doc=str(rng.randint(100000, 999999)))

    for dia in [1, 2, 3, 4, 8, 9, 11, 14, 15, 16, 18, 21, 22, 23, 25, 28, 29, 30]:
        ted_orbe(dia)
    add(1, "TAR PACOTE SERVICOS PJ", "-389.00", "TARIFA")
    add(2, "TRANSF ENTRE CONTAS P/ BCO MERIDIANO AG 0091 CC 90233-5", "-180000.00", "TRANSF_MERIDIANO")
    add(4, "PAGTO SALARIOS - FOLHA 08/2026", "-184732.18", "FOLHA")
    add(4, "BOLETO CONDOMINIO ED. FARIA LIMA 3311", "-28450.00", "CONDOMINIO")
    rep = [x for x in LOJISTAS if x["checkout"]]
    rng.shuffle(rep)
    for i, x in enumerate(rep):
        dia = 8 if i < 8 else 9
        add(dia, "PIX ENVIADO " + nome_aurora(x), str(-round(rng.uniform(4000, 38000), 2)), "REPASSE_LOJISTA", lk=x["k"])
    add(10, "PAGTO ISS SAO PAULO DAMSP 08/2026", "-6877.40", "ISS")
    add(11, "TRANSF ENTRE CONTAS P/ BCO MERIDIANO AG 0091 CC 90233-5", "-210000.00", "TRANSF_MERIDIANO")
    add(15, "DEB AUT AMAZON WEB SERVICES BR", "-41208.77", "CLOUD")
    add(15, "CARTAO PJ - GOOGLE WORKSPACE", "-2210.00", "SOFTWARE")
    add(16, "PIX RECEBIDO 41203887000155 SEM IDENTIFICACAO", "4750.00", "PIX_NAO_IDENTIFICADO")
    add(18, "PIX ENVIADO CONTABILIDADE PRECISA LTDA", "-6800.00", "CONTABILIDADE")
    add(18, "GPS INSS 08/2026", "-51320.66", "INSS")
    add(18, "GRF FGTS 08/2026", "-14778.57", "FGTS")
    add(21, "DEBITO DARF 0561 IRRF", "-22104.35", "IRRF")
    add(21, "TRANSF ENTRE CONTAS P/ BCO MERIDIANO AG 0091 CC 90233-5", "-190000.00", "TRANSF_MERIDIANO")
    rng.shuffle(rep)
    for i, x in enumerate(rep[:9]):
        dia = 22 if i < 4 else 24
        add(dia, "PIX ENVIADO " + nome_aurora(x), str(-round(rng.uniform(4000, 38000), 2)), "REPASSE_LOJISTA", lk=x["k"])
    add(23, "DEVOLUCAO PIX - CAMINHOS INTERC", "8900.00", "DEVOLUCAO_PIX", lk="caminhos")
    add(25, "DEBITO DARF 8109 PIS", "-3811.20", "PIS")
    add(25, "DEBITO DARF 2172 COFINS", "-17582.46", "COFINS")
    add(29, "TAR TED/PIX AVULSAS 09/2026", "-187.50", "TARIFA")
    add(30, "TRANSF ENTRE CONTAS P/ BCO MERIDIANO AG 0091 CC 90233-5", "-150000.00", "TRANSF_MERIDIANO_EM_TRANSITO")
    ordem_cat = {"ORBE_TED": 0}
    ev.sort(key=lambda e: (e["data"], ordem_cat.get(e["cat"], 1)))
    return ev


def escrever_aurora(ev, caminho):
    saldo = D("812334.81")
    linhas = ["Banco Aurora S.A.;;;;", "Extrato de Conta Corrente - Pessoa Jurídica;;;;",
              "Cliente: TARVO PAGAMENTOS LTDA;CNPJ: 84.626.315/0001-39;;;", "Agência: 0318;Conta: 4471-0;;;",
              "Período: 01/09/2026 a 30/09/2026;;;;", ";;;;", "Data;Lançamento;Documento;Valor (R$);Saldo (R$)",
              f"31/08/2026;SALDO ANTERIOR;;;{br(saldo, True)}"]
    for e in ev:
        saldo += e["valor"]
        linhas.append(f"{dbr(e['data'])};{e['desc']};{e['doc']};{br(e['valor'], True)};{br(saldo, True)}")
        e["linha"] = len(linhas)  # numero da linha fisica (1-based)
        e["ref"] = f"L{e['linha']}"
    linhas += [";;;;", f";SALDO FINAL;;;{br(saldo, True)}", ";Os lançamentos estão sujeitos a confirmação.;;;"]
    with open(caminho, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(linhas) + "\n")


# -------------------------------------------------------- extrato Meridiano
def memo_merid(x):
    nome = x["razao"].upper()
    if x["k"] in ("polar",):
        nome = sem_acento(nome)  # o banco as vezes tira o acento...
    else:  # ...e as vezes troca por caractere invalido (U+FFFD), como no OFX real
        nome = "".join(c if ord(c) < 128 else "�" for c in nome)
    return "PIX RECEB " + nome[:24]


def gerar_meridiano(rng):
    ev = []

    def add(dia, valor, memo, cat, lk=None):
        v = D(valor)
        ev.append(dict(data=d(dia), valor=v, memo=memo, cat=cat, lk=lk, tipo="CREDIT" if v > 0 else "DEBIT"))

    rem = [x for x in LOJISTAS if x["remessa"]]
    for dia in [1, 8, 15, 22, 29]:
        for x in rng.sample(rem, rng.randint(3, 5)):
            add(dia, str(round(rng.uniform(9000, 95000), 2)), memo_merid(x), "RECURSO_REMESSA", x["k"])
    for dia in [1, 2, 4, 8, 10, 14, 16, 18, 22, 24, 28, 30]:
        add(dia, str(-round(rng.uniform(60000, 210000), 2)), "TED ENVIADA PONTE CAMBIO SA LIQ OPERACOES", "LIQ_PONTE")
    add(1, "-129.90", "TARIFA MANUTENCAO CONTA PJ", "TARIFA")
    add(2, "180000.00", "TED RECEBIDA TARVO PAGAMENTOS LTDA - MESMA TITULARIDADE", "TRANSF_RECEBIDA")
    add(4, "31942.06", "TED RECEB PONTE CAMBIO SA REPASSE COMISSAO 08/2026", "REPASSE_COMISSAO_PONTE")
    add(5, "-95000.00", "APLICACAO AUTOMATICA CDB DI", "APLICACAO")
    add(9, "96210.00", "RESGATE AUTOMATICO CDB DI", "RESGATE")
    add(11, "210000.00", "TED RECEBIDA TARVO PAGAMENTOS LTDA - MESMA TITULARIDADE", "TRANSF_RECEBIDA")
    add(12, "-1830.44", "SEGURO EMPRESARIAL PARC 09/12", "SEGURO")
    add(21, "190000.00", "TED RECEBIDA TARVO PAGAMENTOS LTDA - MESMA TITULARIDADE", "TRANSF_RECEBIDA")
    add(25, "-120000.00", "APLICACAO AUTOMATICA CDB DI", "APLICACAO")
    add(26, "-48.21", "IOF S/ OPERACAO CREDITO", "IOF")
    add(29, "60817.00", "RESGATE AUTOMATICO CDB DI", "RESGATE")
    add(30, "1912.66", "RENDIMENTO APLIC AUTOMATICA", "RENDIMENTO")
    add(30, "-450.00", "TARIFA PACOTE CAMBIO", "TARIFA")
    ordem = {"CREDIT": 0, "DEBIT": 1}
    ev.sort(key=lambda e: (e["data"], ordem[e["tipo"]], e["cat"] not in ("RECURSO_REMESSA",)))
    n = 1
    for e in ev:
        n += rng.randint(2, 11)
        e["fitid"] = f"MRD202609{n:04d}"
        e["ref"] = e["fitid"]
        e["checknum"] = str(rng.randint(10000, 999999))
    return ev


def escrever_meridiano(ev, caminho):
    cab = """OFXHEADER:100
DATA:OFXSGML
VERSION:102
SECURITY:NONE
ENCODING:USASCII
CHARSET:1252
COMPRESSION:NONE
OLDFILEUID:NONE
NEWFILEUID:NONE

<OFX>
<SIGNONMSGSRSV1>
<SONRS>
<STATUS>
<CODE>0
<SEVERITY>INFO
</STATUS>
<DTSERVER>20261001063000[-3:BRT]
<LANGUAGE>POR
</SONRS>
</SIGNONMSGSRSV1>
<BANKMSGSRSV1>
<STMTTRNRS>
<TRNUID>1001
<STATUS>
<CODE>0
<SEVERITY>INFO
</STATUS>
<STMTRS>
<CURDEF>BRL
<BANKACCTFROM>
<BANKID>0477
<BRANCHID>0091
<ACCTID>90233-5
<ACCTTYPE>CHECKING
</BANKACCTFROM>
<BANKTRANLIST>
<DTSTART>20260901000000[-3:BRT]
<DTEND>20260930235959[-3:BRT]"""
    partes = cab.split("\n")
    for e in ev:
        partes += ["<STMTTRN>", f"<TRNTYPE>{e['tipo']}", f"<DTPOSTED>{e['data']:%Y%m%d}120000[-3:BRT]",
                   f"<TRNAMT>{e['valor']:.2f}", f"<FITID>{e['fitid']}", f"<CHECKNUM>{e['checknum']}",
                   f"<MEMO>{e['memo']}", "</STMTTRN>"]
    saldo = D("905118.42") + sum(e["valor"] for e in ev)
    partes += ["</BANKTRANLIST>", "<LEDGERBAL>", f"<BALAMT>{saldo:.2f}", "<DTASOF>20260930235959[-3:BRT]",
               "</LEDGERBAL>", "</STMTRS>", "</STMTTRNRS>", "</BANKMSGSRSV1>", "</OFX>"]
    with open(caminho, "w", encoding="utf-8", newline="") as f:  # bytes UTF-8 apesar do CHARSET:1252, como o real
        f.write("\r\n".join(partes) + "\r\n")


PLANO_DE_CONTAS = """COD_CONTA;DESCRICAO;TIPO;NATUREZA;ACEITA_LANCTO;COD_SUPERIOR;DT_CRIACAO
1;ATIVO;SINTETICA;D;N;;15/01/2024
1.1;ATIVO CIRCULANTE;SINTETICA;D;N;1;15/01/2024
1.1.01;DISPONIBILIDADES;SINTETICA;D;N;1.1;15/01/2024
1.1.01.001;Banco Aurora C/C 4471-0;ANALITICA;D;S;1.1.01;15/01/2024
1.1.01.002;Banco Meridiano C/C 90233-5 (Câmbio);ANALITICA;D;S;1.1.01;15/01/2024
1.1.02;APLICAÇÕES FINANCEIRAS;SINTETICA;D;N;1.1;15/01/2024
1.1.02.001;Aplicação automática Meridiano;ANALITICA;D;S;1.1.02;15/01/2024
1.1.03;CRÉDITOS A RECEBER;SINTETICA;D;N;1.1;15/01/2024
1.1.03.001;Recebíveis de adquirente - Orbe;ANALITICA;D;S;1.1.03;15/01/2024
1.1.03.002;Comissões de intermediação a receber - câmbio;ANALITICA;D;S;1.1.03;15/01/2024
1.1.03.003;Adiantamentos diversos;ANALITICA;D;S;1.1.03;15/01/2024
1.1.09;CONTAS TRANSITÓRIAS;SINTETICA;D;N;1.1;15/01/2024
1.1.09.001;Transferências entre contas próprias;ANALITICA;D;S;1.1.09;15/01/2024
1.1.09.002;Créditos não identificados;ANALITICA;D;S;1.1.09;15/01/2024
2;PASSIVO;SINTETICA;C;N;;15/01/2024
2.1;PASSIVO CIRCULANTE;SINTETICA;C;N;2;15/01/2024
2.1.01;OBRIGAÇÕES COM LOJISTAS;SINTETICA;C;N;2.1;15/01/2024
2.1.01.001;Repasses a lojistas a pagar;ANALITICA;C;S;2.1.01;15/01/2024
2.1.01.002;Devoluções e créditos a lojistas;ANALITICA;C;S;2.1.01;15/01/2024
2.1.02;RECURSOS DE CLIENTES - REMESSAS;SINTETICA;C;N;2.1;15/01/2024
2.1.02.001;Recursos de lojistas para remessas a liquidar;ANALITICA;C;S;2.1.02;15/01/2024
2.1.03;OBRIGAÇÕES FISCAIS;SINTETICA;C;N;2.1;15/01/2024
2.1.03.001;PIS a recolher;ANALITICA;C;S;2.1.03;15/01/2024
2.1.03.002;COFINS a recolher;ANALITICA;C;S;2.1.03;15/01/2024
2.1.03.003;ISS a recolher;ANALITICA;C;S;2.1.03;15/01/2024
2.1.03.004;IRRF a recolher;ANALITICA;C;S;2.1.03;15/01/2024
2.1.04;OBRIGAÇÕES TRABALHISTAS;SINTETICA;C;N;2.1;15/01/2024
2.1.04.001;Salários a pagar;ANALITICA;C;S;2.1.04;15/01/2024
2.1.04.002;INSS a recolher;ANALITICA;C;S;2.1.04;15/01/2024
2.1.04.003;FGTS a recolher;ANALITICA;C;S;2.1.04;15/01/2024
2.1.05;FORNECEDORES;SINTETICA;C;N;2.1;15/01/2024
2.1.05.001;Fornecedores nacionais;ANALITICA;C;S;2.1.05;15/01/2024
3;RECEITAS;SINTETICA;C;N;;15/01/2024
3.1;RECEITA OPERACIONAL;SINTETICA;C;N;3;15/01/2024
3.1.01;Receita de MDR - Checkout;ANALITICA;C;S;3.1;15/01/2024
3.1.02;Receita de antecipação;ANALITICA;C;S;3.1;15/01/2024
3.1.03;Receita de intermediação de câmbio;ANALITICA;C;S;3.1;15/01/2024
3.2;RECEITAS FINANCEIRAS;SINTETICA;C;N;3;15/01/2024
3.2.01;Rendimentos de aplicações;ANALITICA;C;S;3.2;15/01/2024
4;DESPESAS;SINTETICA;D;N;;15/01/2024
4.1;DESPESAS COM PESSOAL;SINTETICA;D;N;4;15/01/2024
4.1.01;Salários e ordenados;ANALITICA;D;S;4.1;15/01/2024
4.2;DESPESAS ADMINISTRATIVAS;SINTETICA;D;N;4;15/01/2024
4.2.01;Tarifas bancárias;ANALITICA;D;S;4.2;15/01/2024
4.2.02;Aluguel e condomínio;ANALITICA;D;S;4.2;15/01/2024
4.2.03;Cloud e infraestrutura;ANALITICA;D;S;4.2;15/01/2024
4.2.04;Softwares e assinaturas;ANALITICA;D;S;4.2;15/01/2024
4.2.05;Seguros;ANALITICA;D;S;4.2;15/01/2024
4.2.06;Serviços contábeis e jurídicos;ANALITICA;D;S;4.2;15/01/2024
4.3;DESPESAS FINANCEIRAS;SINTETICA;D;N;4;15/01/2024
4.3.01;IOF;ANALITICA;D;S;4.3;15/01/2024
4.3.02;Juros e multas;ANALITICA;D;S;4.3;15/01/2024"""


def escrever_plano(caminho):
    with open(caminho, "w", encoding="cp1252", newline="") as f:
        f.write(PLANO_DE_CONTAS.replace("\n", "\r\n") + "\r\n")


NOMES = dict(
    orbe01="01_orbe_relatorio_comissoes_set2026.csv",
    tabela="02_tabela_taxas_contratadas.xlsx",
    eml="03_email_ponte_cambio_relatorio_set2026.eml",
    plano="04_plano_de_contas_fenix_erp.csv",
    aurora="05a_extrato_banco_aurora_cc4471_set2026.csv",
    meridiano="05b_extrato_banco_meridiano_cc90233_set2026.ofx",
    orbe06="06_orbe_ajustes_tardios_set2026.csv",
)


def construir_mundo():
    rng = random.Random(SEED)
    montar_cadastro(rng)
    orbe = gerar_orbe(rng)
    ponte = gerar_ponte(rng)
    aurora = gerar_aurora(rng)
    merid = gerar_meridiano(rng)
    return rng, dict(orbe=orbe, ponte=ponte, aurora=aurora, merid=merid)


def escrever_entradas(rng, mundo, pasta):
    os.makedirs(pasta, exist_ok=True)
    escrever_orbe(mundo["orbe"], os.path.join(pasta, NOMES["orbe01"]), "01", rng)
    escrever_orbe(mundo["orbe"], os.path.join(pasta, NOMES["orbe06"]), "06", rng)
    mundo["linhas_tab"] = escrever_tabela(os.path.join(pasta, NOMES["tabela"]))
    xbytes = escrever_ponte_xlsx(mundo["ponte"], None)
    escrever_eml(xbytes, os.path.join(pasta, NOMES["eml"]))
    escrever_plano(os.path.join(pasta, NOMES["plano"]))
    escrever_aurora(mundo["aurora"], os.path.join(pasta, NOMES["aurora"]))
    escrever_meridiano(mundo["merid"], os.path.join(pasta, NOMES["meridiano"]))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--saida", default=AQUI)
    ap.add_argument("--so-entradas", action="store_true")
    a = ap.parse_args()
    rng, mundo = construir_mundo()
    escrever_entradas(rng, mundo, os.path.join(a.saida, "entrada"))
    if not a.so_entradas:
        import sys
        sys.path.insert(0, AQUI)
        sys.modules.setdefault("gerar_massa", sys.modules[__name__])  # gabarito usa o MESMO cadastro montado aqui
        import gabarito  # noqa: E402  (fica em gabarito.py, ao lado deste arquivo)
        gabarito.gerar(mundo, a.saida)
    print("ok:", os.path.abspath(a.saida))


if __name__ == "__main__":
    main()
